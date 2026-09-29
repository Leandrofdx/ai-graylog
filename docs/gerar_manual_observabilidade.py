#!/usr/bin/env python3
"""Gera o Manual de Observabilidade (DOCX) — documento autônomo; lab só como exemplo."""
from __future__ import annotations

from datetime import date
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

OUT = Path(__file__).resolve().parent / "Manual_Observabilidade.docx"

# Paleta
INK = RGBColor(0x0F, 0x17, 0x2A)
MUTED = RGBColor(0x55, 0x55, 0x65)
ACCENT = RGBColor(0x1E, 0x3A, 0x5F)
RULE = RGBColor(0xC5, 0xCB, 0xD3)


def set_run_font(run, *, bold=False, size=11, color=None, italic=False, name="Calibri"):
    run.bold = bold
    run.italic = italic
    run.font.size = Pt(size)
    run.font.name = name
    r = run._element
    rPr = r.get_or_add_rPr()
    rFonts = rPr.get_or_add_rFonts()
    rFonts.set(qn("w:ascii"), name)
    rFonts.set(qn("w:hAnsi"), name)
    rFonts.set(qn("w:eastAsia"), name)
    if color:
        run.font.color.rgb = color


def add_heading(doc, text, level=1):
    p = doc.add_heading(text, level=level)
    sizes = {1: 16, 2: 13, 3: 12}
    for run in p.runs:
        set_run_font(run, bold=True, size=sizes.get(level, 12), color=INK if level == 1 else ACCENT)
    return p


def add_para(doc, text, *, bold=False, italic=False, size=11, space_after=8, color=None):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
    run = p.add_run(text)
    set_run_font(run, bold=bold, italic=italic, size=size, color=color or INK)
    return p


def add_bullet(doc, text, *, level=0):
    p = doc.add_paragraph(text, style="List Bullet")
    p.paragraph_format.left_indent = Cm(0.5 + 0.5 * level)
    p.paragraph_format.space_after = Pt(3)
    for run in p.runs:
        set_run_font(run, size=11, color=INK)
    return p


def add_numbered(doc, text):
    p = doc.add_paragraph(text, style="List Number")
    p.paragraph_format.space_after = Pt(3)
    for run in p.runs:
        set_run_font(run, size=11, color=INK)
    return p


def add_callout(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.space_after = Pt(10)
    p.paragraph_format.left_indent = Cm(0.4)
    run = p.add_run(text)
    set_run_font(run, italic=True, size=10, color=MUTED)
    return p


def add_table(doc, headers, rows):
    table = doc.add_table(rows=1 + len(rows), cols=len(headers))
    table.style = "Table Grid"
    hdr = table.rows[0].cells
    for i, h in enumerate(headers):
        hdr[i].text = ""
        p = hdr[i].paragraphs[0]
        run = p.add_run(h)
        set_run_font(run, bold=True, size=9, color=INK)
    for r_i, row in enumerate(rows):
        cells = table.rows[r_i + 1].cells
        for c_i, val in enumerate(row):
            cells[c_i].text = ""
            p = cells[c_i].paragraphs[0]
            run = p.add_run(str(val))
            set_run_font(run, size=9, color=INK)
    doc.add_paragraph()
    return table


def add_code(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(10)
    p.paragraph_format.left_indent = Cm(0.35)
    run = p.add_run(text)
    set_run_font(run, size=8, name="Consolas", color=ACCENT)
    return p


def add_hr_note(doc, text):
    add_para(doc, text, italic=True, size=9, color=MUTED, space_after=12)


def build():
    doc = Document()
    section = doc.sections[0]
    section.top_margin = Cm(2.2)
    section.bottom_margin = Cm(2.2)
    section.left_margin = Cm(2.3)
    section.right_margin = Cm(2.3)

    # ── Capa ──
    for _ in range(2):
        doc.add_paragraph()
    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = title.add_run("MANUAL DE OBSERVABILIDADE")
    set_run_font(r, bold=True, size=26, color=INK)

    sub = doc.add_paragraph()
    sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = sub.add_run("Logs · Traces · Métricas")
    set_run_font(r, size=16, color=ACCENT)

    sub2 = doc.add_paragraph()
    sub2.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = sub2.add_run(
        "Mentalidade, contrato de telemetria, implementação\n"
        "e playbook de investigação — documento autônomo"
    )
    set_run_font(r, size=12, color=MUTED)

    meta = doc.add_paragraph()
    meta.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = meta.add_run(
        f"\n\nVersão consolidada · {date.today().strftime('%d/%m/%Y')}\n"
        "Público: engenharia, SRE, produto e líderes técnicos\n"
        "Exemplos de prática: lab Assistente de Vendas (opcional)"
    )
    set_run_font(r, size=10, italic=True, color=MUTED)

    note = doc.add_paragraph()
    note.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = note.add_run(
        "\n\nComo ler este manual\n"
        "O corpo do texto é genérico e vale em qualquer stack (Datadog, Elastic, Grafana,\n"
        "CloudWatch, Jaeger, etc.). Blocos marcados como «Exemplo (lab)» são ilustração\n"
        "do Assistente de Vendas / Graylog / Jaeger — treino, não pré-requisito."
    )
    set_run_font(r, size=10, italic=True, color=MUTED)

    guide = doc.add_paragraph()
    guide.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = guide.add_run(
        "\n\nFrase-guia: métrica acusa → trace localiza → log explica."
    )
    set_run_font(r, bold=True, size=11, color=ACCENT)

    doc.add_page_break()

    # ── Sumário ──
    add_heading(doc, "Sumário", 1)
    for item in [
        "1. Objetivo e escopo",
        "2. Mentalidade",
        "3. O problema que a telemetria deve separar",
        "4. Os três pilares (logs, métricas, traces)",
        "5. Trace e span",
        "6. Glossário essencial",
        "7. Correlação — o coração do método",
        "8. Exemplos de código (padrão + lab)",
        "9. Como achar o trace_id (fluxo profissional)",
        "10. Responsabilidades das ferramentas",
        "11. Playbook de incidente",
        "12. Exemplos guiados de diagnóstico",
        "13. Análise passo a passo (treino com lab)",
        "14. Contrato de telemetria",
        "15. Implementação do zero",
        "16. Dashboards com intenção",
        "17. Do laboratório ao sistema real",
        "18. Maturidade",
        "19. Anti-padrões",
        "20. Critério de prontidão",
        "Anexo A — Queries úteis",
        "Anexo B — Arquitetura de referência (lab)",
        "Anexo C — URLs e credenciais de treino",
    ]:
        add_para(doc, item, size=10, space_after=2)

    doc.add_page_break()

    # ── 1 ──
    add_heading(doc, "1. Objetivo e escopo", 1)
    add_para(
        doc,
        "Este manual capacita o profissional a: (i) entender observabilidade como perguntas "
        "com evidência, não como “ter ferramenta”; (ii) modelar a jornada de negócio e um "
        "contrato de telemetria; (iii) implementar correlação, erros tipados e eventos de "
        "negócio no código; (iv) investigar incidentes na ordem certa (classificar → localizar "
        "→ tendência); (v) levar o mesmo método ao sistema real da empresa.",
    )
    add_para(
        doc,
        "Observabilidade é a capacidade de, a partir de sinais emitidos pelo sistema, "
        "reconstruir o que ocorreu em uma tentativa específica e responder perguntas novas — "
        "por exemplo: foi estoque ou crédito? em qual etapa? foi regra de negócio rápida ou "
        "gargalo de banco? caso isolado ou degradação do domínio?",
    )
    add_callout(
        doc,
        "O laboratório Assistente de Vendas (Graylog + Jaeger + OpenTelemetry) é ambiente de "
        "prática. Os conceitos, o playbook e os contratos valem com qualquer backend equivalente.",
    )

    # ── 2 ──
    add_heading(doc, "2. Mentalidade", 1)
    add_para(doc, "Mudança de mentalidade necessária:", bold=True)
    add_table(
        doc,
        ["Antes (comum)", "Depois (maduro)"],
        [
            ["“Olha o log do servidor”", "“Qual error_type e em qual etapa da jornada?”"],
            ["“A API deu 409”", "“409 InsufficientStock no SKU, qty=…, stock=…”"],
            ["“Está lento”", "“p95 da operação X; span SQL Y ms no trace T”"],
            ["Culpa genérica no “backend”", "Domínio: estoque vs crédito vs pagamento vs DB"],
            ["Dashboard por vaidade", "Dashboard que responde uma pergunta de negócio"],
            ["Ferramenta primeiro", "Jornada + contrato de telemetria primeiro"],
        ],
    )
    add_para(
        doc,
        "Princípio central: instrumentar a cadeia da operação de negócio "
        "(headers + trace_id + error_type + biz_event), não “o servidor em geral”.",
        italic=True,
    )

    # ── 3 ──
    add_heading(doc, "3. O problema que a telemetria deve separar", 1)
    add_para(
        doc,
        "Quando a operação falha, o relato típico é vago (“travou na pré-venda”). "
        "Sem correlação e sem erros tipados, o time discute opinião. Com o modelo deste manual, "
        "o time discute evidência: etapa, API, causa de negócio ou gargalo técnico.",
    )
    add_table(
        doc,
        ["Classe", "Exemplo", "Sinal principal"],
        [
            ["Negócio", "Sem estoque / crédito negado", "error_type + campos (item_id, cpf, amount)"],
            ["Jornada", "Parou depois de simular plano", "journey_step / biz_event"],
            ["Performance", "Demora demais", "duration_ms, p95, span longo"],
            ["Infra / DB", "Timeout de pool sob carga", "DbPoolTimeout, span DB, RED"],
            ["Cliente / sessão", "Clique / ticket de suporte", "request_id / ID de sessão"],
        ],
    )

    # ── 4 ──
    add_heading(doc, "4. Os três pilares", 1)

    add_heading(doc, "4.1 Log", 2)
    add_para(
        doc,
        "Registro estruturado de um acontecimento: mensagem + campos (status HTTP, caminho da API, "
        "duração, tipo de erro, identificadores de negócio).",
    )
    add_para(doc, "Analogia: o livro de ocorrências da loja — “às 14:32 a pré-venda do SKU-88 foi recusada por estoque insuficiente”.", italic=True)
    add_para(doc, "Pergunta que responde: o que aconteceu e com quais atributos de negócio?")
    add_hr_note(doc, "Exemplo (lab): logs GELF → Graylog. Em outros ambientes: JSON → ELK, Loki, Datadog Logs, CloudWatch.")

    add_heading(doc, "4.2 Métrica", 2)
    add_para(
        doc,
        "Série numérica ao longo do tempo: volume de requisições, taxa de erro, latência "
        "(média, máximos, percentis).",
    )
    add_para(doc, "Analogia: o placar eletrônico — quantas vendas por minuto, quantos erros, se o tempo subiu no pico.", italic=True)
    add_para(doc, "Pergunta: está piorando? em qual serviço/domínio? é tendência ou pico?")
    add_para(doc, "Modelo mínimo de serviço: RED — Rate (vazão), Errors (erros), Duration (latência).")
    add_hr_note(doc, "Exemplo (lab): spanmetrics no Collector → Prometheus → Jaeger Monitor (SPM).")

    add_heading(doc, "4.3 Trace", 2)
    add_para(
        doc,
        "Registro do caminho de uma requisição, do ingresso HTTP até dependências internas "
        "(handlers, consultas ao banco, chamadas a outros serviços), com tempos parciais.",
    )
    add_para(doc, "Analogia: o filme completo daquela tentativa de venda.", italic=True)
    add_para(doc, "Pergunta: onde, dentro dessa chamada, o tempo foi gasto ou a falha ocorreu?")
    add_hr_note(doc, "Exemplo (lab): OpenTelemetry → Collector → Jaeger.")
    add_para(
        doc,
        "Os três pilares se complementam: log classifica o fato; métrica mostra o padrão; "
        "trace localiza o trecho interno.",
        bold=True,
    )

    # ── 5 ──
    add_heading(doc, "5. Trace e span", 1)
    add_heading(doc, "5.1 Trace", 2)
    add_para(
        doc,
        "Um trace representa o ciclo de vida de uma requisição. Possui um identificador global: "
        "trace_id. Com ele é possível localizar a mesma tentativa no backend de logs e no de traces. "
        "Esse vínculo é o núcleo da correlação.",
    )
    add_heading(doc, "5.2 Span", 2)
    add_para(
        doc,
        "Um span é uma unidade de trabalho dentro do trace: intervalo com nome, início, duração, "
        "status (sucesso/erro) e atributos opcionais (error.type, http.status_code, journey.step, etc.).",
    )
    add_para(doc, "Analogia: trace = a corrida completa; span = cada trecho (largada, curva, reta final).", italic=True)
    add_code(
        doc,
        "Trace — CreatePreSales\n"
        "└─ Span HTTP POST /SalesOrder/api/CreatePreSales\n"
        "    ├─ Span SELECT customer\n"
        "    ├─ Span SELECT product FOR UPDATE (estoque)\n"
        "    └─ (falha InsufficientStock — venda não gravada)",
    )
    add_para(
        doc,
        "O span não executa a regra de negócio; ele mede e anota um pedaço da execução. "
        "Sem spans, sabe-se apenas que “a API demorou 2 s”. Com spans, distingue-se "
        "“1,8 s no SQL” de “validação de estoque em 3 ms”.",
    )

    # ── 6 ──
    add_heading(doc, "6. Glossário essencial", 1)
    add_table(
        doc,
        ["Termo", "Definição"],
        [
            ["Observabilidade", "Inferir estado interno por sinais externos, com perguntas novas"],
            ["Log / Métrica / Trace", "Fato · tendência · caminho de uma request"],
            ["Span", "Unidade de trabalho dentro do trace"],
            ["trace_id / span_id", "ID do filme / ID da cena"],
            ["traceparent", "Header W3C Trace Context (padrão de produção)"],
            ["request_id", "ID estável da chamada (suporte/gateway)"],
            ["error_type", "Classe estável da falha (InsufficientStock, DbTimeout)"],
            ["biz_event", "Fato de negócio (sale_created, proposal_integrated)"],
            ["journey_step", "Etapa nomeada da jornada (sale.create_presale)"],
            ["OpenTelemetry (OTel)", "Padrão aberto de instrumentação; instrumentar uma vez, trocar backend"],
            ["OTLP", "Protocolo de exportação OpenTelemetry"],
            ["Collector", "Pipeline: recebe, processa e exporta telemetria"],
            ["GELF", "Formato de log estruturado (comum com Graylog); equivalente: JSON"],
            ["RED", "Rate · Errors · Duration"],
            ["SPM / APM", "Visão de performance por serviço/operação"],
            ["Spanmetrics", "Derivação de métricas RED a partir de spans"],
            ["p95", "95% das chamadas mais rápidas que esse valor"],
            ["GMV", "Soma de amount das vendas"],
            ["Bounded context", "Fronteira de domínio (venda, estoque, pagamento…)"],
        ],
    )

    # ── 7 ──
    add_heading(doc, "7. Correlação — o coração do método", 1)
    add_para(
        doc,
        "Correlacionar = garantir que a mesma tentativa carregue IDs iguais (ou ligáveis) "
        "no cliente, no log e no trace. Sem o mesmo identificador nos três sinais, há ruído. "
        "Com correlação, há uma narrativa por tentativa.",
    )

    add_heading(doc, "7.1 Identificadores", 2)
    add_table(
        doc,
        ["Identificador", "Escopo", "Uso"],
        [
            ["X-Request-Id / request_id", "Uma chamada HTTP", "Linha a linha / ticket / suporte"],
            ["ID de sessão / X-Test-Run-Id", "Sessão UI ou run de carga", "Jornada completa (várias calls)"],
            ["X-Journey-Step", "Etapa nomeada", "Funil e localização semântica"],
            ["trace_id", "Trace OTel/APM", "Amarrar log ↔ trace"],
            ["error_type", "Classe de falha", "Dashboard e alerta de negócio"],
            ["biz_event", "Fato de negócio", "Funil / GMV (≠ status HTTP)"],
            ["IDs de negócio", "Pedido, proposta, etc.", "Entrada a partir de CRM/suporte"],
            ["staff_id / user_id", "Operador", "“Quem” (respeitando PII)"],
        ],
    )

    add_heading(doc, "7.2 Quem gera / quem consome", 2)
    add_table(
        doc,
        ["ID", "Quem gera", "Quem propaga", "Quem usa"],
        [
            ["request_id", "Cliente ou API/gateway", "Header em cada call", "Uma chamada no log"],
            ["ID de sessão", "Cliente no início da sessão", "Todas as calls da sessão", "Juntar a jornada"],
            ["journey_step", "Cliente/BFF conforme a ação", "Header naquela call", "Etapa sem depender da URL"],
            ["trace_id", "OTel/APM na API (automático)", "Log + spans + header de resposta", "Mesmo filme em log e APM"],
            ["error_type", "API no ramo de falha", "Log + atributo do span", "Alertas tipados"],
            ["biz_event", "API no sucesso de negócio", "Log (+ attrs de span)", "Funil / GMV"],
        ],
    )

    add_heading(doc, "7.3 Analogia da comanda", 2)
    add_bullet(doc, "ID de sessão = número da comanda na mesa")
    add_bullet(doc, "request_id = cada item lançado na comanda")
    add_bullet(doc, "journey_step = entrada / prato / sobremesa")
    add_bullet(doc, "trace_id = código da filmagem daquela comanda na cozinha")
    add_bullet(doc, "error_type = “faltou ingrediente” (padronizado), não só “pedido recusado”")

    add_heading(doc, "7.4 Onde aplicar no código", 2)
    add_numbered(doc, "Middleware de entrada — ler headers; gerar request_id se faltar; taguear span.")
    add_numbered(doc, "Logger — enricher com trace_id / span_id em todo record.")
    add_numbered(doc, "Middleware de saída — devolver traceparent ou X-Trace-Id + X-Request-Id; log de envelope HTTP.")
    add_numbered(doc, "Regras de negócio — error_type + campos da regra; no sucesso, biz_event + IDs.")
    add_numbered(doc, "Clientes HTTP/gRPC — propagação automática do contexto.")
    add_numbered(doc, "Filas/async — traceparent / trace_id na mensagem.")
    add_para(
        doc,
        "Não se “configura correlação só na ferramenta”: a ferramenta recebe o que o código emitiu.",
        italic=True,
    )

    add_heading(doc, "7.5 Fluxo numa funcionalidade (ex.: pré-venda)", 2)
    add_code(
        doc,
        "[1] Cliente  → headers (Request-Id, sessão, Journey-Step, Staff-Id)\n"
        "[2] Cliente  → POST …/CreatePreSales { … }\n"
        "[3] API      → OTel cria TRACE + SPAN HTTP\n"
        "[4] API      → regra; se falhar: error_type + log + span ERROR\n"
        "[5] API      → se ok: biz_event + IDs de negócio\n"
        "[6] API      → response: X-Trace-Id / traceparent (+ opcional deep links ops)\n"
        "[7] Ops      → a partir do header ou do log, abre APM e busca de logs",
    )
    add_callout(
        doc,
        "Sobre “Rastreio” na UI: útil em lab ou backoffice de suporte. Não é requisito do método "
        "nem algo para colocar no app do vendedor. Em produção o contrato são headers + logs + APM (§9).",
    )

    # ── 8 ──
    add_heading(doc, "8. Exemplos de código (padrão + lab)", 1)
    add_para(doc, "Os trechos mostram onde cada peça vive. Adapte nomes ao seu framework.")

    add_heading(doc, "8.1 Cliente — gerar e enviar correlação", 2)
    add_code(
        doc,
        "function ensureSessionId() { /* uma vez por aba/sessão */ }\n"
        "function newRequestId() {\n"
        "  if (globalThis.crypto?.randomUUID) return crypto.randomUUID();\n"
        "  return `req-${Date.now().toString(36)}-${Math.random().toString(36).slice(2, 10)}`;\n"
        "}\n\n"
        "function headers(step = \"\") {\n"
        "  return {\n"
        "    \"X-Request-Id\": newRequestId(),\n"
        "    \"X-Test-Run-Id\": ensureSessionId(),  // ou omitir fora de lab/teste\n"
        "    \"X-Journey-Step\": step,              // ex.: sale.create_presale\n"
        "  };\n"
        "}\n\n"
        "await fetch(\"/SalesOrder/api/CreatePreSales\", {\n"
        "  method: \"POST\",\n"
        "  headers: { \"Content-Type\": \"application/json\", ...headers(\"sale.create_presale\") },\n"
        "  body: JSON.stringify({ /* … */ }),\n"
        "});",
    )

    add_heading(doc, "8.2 API — middleware de entrada", 2)
    add_code(
        doc,
        "@app.before_request\n"
        "def _before():\n"
        "    g.started_at = time.perf_counter()\n"
        "    g.request_id = request.headers.get(\"X-Request-Id\", str(uuid.uuid4()))\n"
        "    g.test_run_id = request.headers.get(\"X-Test-Run-Id\", \"\")\n"
        "    g.journey_step = request.headers.get(\"X-Journey-Step\", \"\")\n"
        "    g.staff_id = request.headers.get(\"X-Staff-Id\", \"\")\n"
        "    _tag_span(**{\n"
        "        \"request.id\": g.request_id,\n"
        "        \"test_run.id\": g.test_run_id,\n"
        "        \"journey.step\": g.journey_step,\n"
        "        \"staff.id\": g.staff_id,\n"
        "    })",
    )

    add_heading(doc, "8.3 API — middleware de saída (contrato de headers)", 2)
    add_code(
        doc,
        "@app.after_request\n"
        "def _after(response):\n"
        "    tid, _ = _ids()  # do span OTel atual\n"
        "    response.headers[\"X-Request-Id\"] = g.request_id\n"
        "    if tid:\n"
        "        response.headers[\"X-Trace-Id\"] = tid\n"
        "        # Em produção: preferir também traceparent (W3C)\n"
        "    # access log: method, route, status, duration_ms, error_type, journey_step, trace_id\n"
        "    return response",
    )

    add_heading(doc, "8.4 Erro tipado (estoque na pré-venda)", 2)
    add_code(
        doc,
        "def _problem(error_type, status, body=None, **fields):\n"
        "    g.error_type = error_type\n"
        "    _tag_span(**{\"error.type\": error_type, \"error\": True})\n"
        "    logger.warning(\"%s status=%s\", error_type, status,\n"
        "                   extra=_extra(error_type=error_type, http_status=status, **fields))\n"
        "    return jsonify({\"error\": error_type, **(body or {})}), status\n\n"
        "if product[\"stock\"] < qty:\n"
        "    return _problem(\n"
        "        \"InsufficientStock\", 409,\n"
        "        {\"available\": product[\"stock\"]},\n"
        "        item_id=item_id, qty=qty, stock=product[\"stock\"],\n"
        "    )",
    )
    add_para(doc, "HTTP 409 sozinho não basta: o nome estável é o que dashboard e alerta usam.", italic=True)

    add_heading(doc, "8.5 Sucesso de negócio", 2)
    add_code(
        doc,
        "logger.info(\"pre-sale created id=%s …\", sale_id,\n"
        "    extra=_extra(\n"
        "        biz_event=\"sale_created\",\n"
        "        sale_id=sale_id, amount=amount, item_id=item_id,\n"
        "        payment_method=…, proposal_id=…,\n"
        "    ))\n"
        "_tag_span(**{\"sale.id\": sale_id, \"sale.amount\": amount, \"item.id\": item_id})",
    )

    add_heading(doc, "8.6 Enricher de trace_id no log + envelope comum", 2)
    add_para(
        doc,
        "No Filter do logging, ler o span ativo OTel e setar record.trace_id / span_id. "
        "Campos base recomendados em todo evento (_extra): serviço/domínio, http_method, "
        "http_path, http_route (template), request_id, journey_step, trace_id, span_id, "
        "staff_id, duration_ms + kwargs (error_type / biz_event / IDs). "
        "A rota template (/orders/{id}) é melhor para dashboard do que a URL crua com ids.",
    )

    # ── 9 ──
    add_heading(doc, "9. Como achar o trace_id (fluxo profissional)", 1)
    add_para(
        doc,
        "Em aplicação real não há “cantinho de Rastreio” no PDV. O ID se obtém pelos canais "
        "de operação que já existem: headers, logs, APM e o ticket/incidente.",
    )
    add_code(
        doc,
        "Sintoma\n"
        "  ├─ alerta de serviço/operação  → APM/SPM → abrir um trace com erro\n"
        "  ├─ horário + endpoint/erro     → Log: error_type / status / path\n"
        "  ├─ request_id (gateway/ticket) → Log: request_id → ler trace_id\n"
        "  └─ ID de negócio / user+tempo  → Log → ler trace_id\n\n"
        "Com trace_id → logs (negócio) + APM (caminho) + RED (tendência)",
    )
    add_table(
        doc,
        ["Fonte", "Onde olhar"],
        [
            ["Header da resposta", "DevTools Network / Postman / curl -i → X-Trace-Id ou traceparent"],
            ["Campo no log", "Evento no indexador → campo trace_id"],
            ["Detalhe do APM", "URL /trace/{id} ou painel do span raiz"],
            ["Ticket", "request_id ou order_id → resolve no log"],
        ],
    )
    add_para(
        doc,
        "O app deve devolver ID correlacionável em toda resposta (incluindo 4xx/5xx). "
        "Isso é contrato de API, não feature de UI.",
        bold=True,
    )
    add_para(
        doc,
        "Se o usuário só disse “deu erro de estoque às 14h”, entre por error_type + janela de tempo — "
        "não por trace_id. O trace_id aparece no evento e aí você aprofunda.",
        italic=True,
    )

    # ── 10 ──
    add_heading(doc, "10. Responsabilidades das ferramentas", 1)
    add_para(doc, "Usar a ferramenta errada é a causa mais comum de investigação longa e inconclusiva.")
    add_table(
        doc,
        ["Papel", "Genérico", "Exemplo (lab)"],
        [
            ["O que aconteceu (negócio)", "Logs", "Graylog"],
            ["Onde no caminho", "Traces", "Jaeger"],
            ["Isolado vs epidemia", "RED / SPM / APM service page", "Jaeger Monitor"],
            ["Pipeline de telemetria", "Collector / agent", "OTel Collector"],
            ["Estado que permite falha real", "DB / deps", "Postgres (estoque/crédito)"],
        ],
    )
    add_heading(doc, "10.1 O que observar", 2)
    add_bullet(doc, "Logs: error_type, http_status, path/route, duration_ms, biz_event, IDs, trace_id")
    add_bullet(doc, "Trace: árvore, span ERROR, atributos, filhos longos (SQL/deps)")
    add_bullet(doc, "RED: rate, error rate, p95 do serviço/operação no intervalo")
    add_heading(doc, "10.2 Ordem padrão", 2)
    add_code(
        doc,
        "Sintoma\n"
        "  → Logs (classificar error_type / biz_event + contexto)\n"
        "  → Trace (localizar span / latência)\n"
        "  → RED/SPM (tendência)\n"
        "  → Ação com evidência citada",
    )
    add_para(
        doc,
        "Não comece por métricas se já tem um caso concreto; não fique só no log se a pergunta "
        "é “onde gastou tempo”.",
        italic=True,
    )

    # ── 11 ──
    add_heading(doc, "11. Playbook de incidente", 1)
    add_numbered(doc, "Sintoma — alerta, usuário, status.")
    add_numbered(doc, "Recorte — serviço/operação + janela de tempo.")
    add_numbered(doc, "Amostra — 1–3 eventos ou traces (não milhares).")
    add_numbered(doc, "Classificar — domínio tipado vs infra.")
    add_numbered(doc, "Contexto — campos de negócio no log.")
    add_numbered(doc, "Caminho — span culpado / gargalo no trace.")
    add_numbered(doc, "Tendência — RED isolado vs subindo.")
    add_numbered(doc, "Agir — código, dados ou infra.")
    add_numbered(doc, "Validar — mesma query/alerta caindo.")
    add_para(
        doc,
        "Essa sequência força classificação antes de otimização — e é o que permite expor "
        "exatamente onde está o problema.",
        italic=True,
    )

    # ── 12 ──
    add_heading(doc, "12. Exemplos guiados de diagnóstico", 1)

    add_heading(doc, "12.1 Pré-venda recusada (estoque)", 2)
    add_table(
        doc,
        ["Etapa", "Evidência", "Interpretação"],
        [
            ["Entrada", "etapa sale.create_presale, HTTP 409", "Falha na criação"],
            ["Log", "error_type=InsufficientStock, item_id, qty, stock", "Causa de negócio"],
            ["Trace", "Span HTTP ERROR; SQL curto", "Regra rápida, não timeout"],
            ["RED", "error rate do domínio venda", "Isolado vs onda"],
        ],
    )
    add_para(
        doc,
        "Veredito: inventário/quantidade — não “escalar Postgres” sem evidência.",
        bold=True,
    )

    add_heading(doc, "12.2 Lentidão com HTTP 200", 2)
    add_table(
        doc,
        ["Etapa", "Evidência", "Interpretação"],
        [
            ["Entrada", "etapa simulação, 200, demora percebida", "Sucesso lento"],
            ["Log", "duration_ms alto", "Quantifica"],
            ["Trace", "Span filho domina o tempo", "Local do custo"],
            ["RED", "p95 do domínio sobe", "Degradação"],
        ],
    )

    add_heading(doc, "12.3 Crédito inconsistente", 2)
    add_table(
        doc,
        ["Etapa", "Evidência", "Interpretação"],
        [
            ["Log", "error_type=CreditNotApproved, product_type, cpf", "Falha antes da pré-venda"],
            ["Funil", "sale.check_credit aparece; create_presale some", "Jornada interrompida"],
            ["RED", "pico no domínio cliente, não venda", "Fronteira correta"],
        ],
    )
    add_para(doc, "Veredito: limites/crédito, não CreatePreSales.", bold=True)

    # ── 13 ──
    add_heading(doc, "13. Análise passo a passo (treino com lab)", 1)
    add_callout(
        doc,
        "Esta seção usa dados reais do laboratório Assistente de Vendas para treinar o olho. "
        "Na sua empresa os nomes de campo e URLs mudam; o método não.",
    )

    add_heading(doc, "13.1 Caso InsufficientStock", 2)
    add_table(
        doc,
        ["Campo", "Valor capturado"],
        [
            ["Operação", "POST /SalesOrder/api/CreatePreSales"],
            ["Status", "409"],
            ["error_type", "InsufficientStock"],
            ["Contexto", "item_id=SKU-42, qty=2000000, stock=999999"],
            ["trace_id", "b050b7b79c6154d1ca2169b9ca7aa7c0"],
            ["journey_step", "sale.create_presale"],
            ["Duração", "~6–9 ms (falha rápida de regra)"],
        ],
    )

    add_heading(doc, "13.2 Entrada profissional (sem dock)", 2)
    add_code(doc, "error_type:InsufficientStock")
    add_para(doc, "Abra o evento → copie trace_id → refine:")
    add_code(doc, "trace_id:b050b7b79c6154d1ca2169b9ca7aa7c0")
    add_para(
        doc,
        "Mensagens típicas: (1) domínio com qty/stock; (2) envelope POST … -> 409. "
        "Conclusão: regra de estoque; não timeout de DB.",
    )

    add_heading(doc, "13.3 Trace no Jaeger", 2)
    add_code(
        doc,
        "POST /SalesOrder/api/CreatePreSales   ~8.8 ms   status=409  error=true\n"
        "  ├─ connect\n"
        "  └─ SELECT lab  (curtos)",
    )
    add_para(doc, "Parou em CreatePreSales; SQL rápido → rejeição de negócio após ler estoque.")

    add_heading(doc, "13.4 SPM / RED", 2)
    add_para(
        doc,
        "Ponto isolado vs taxa subindo no serviço da operação (assistente-venda). "
        "SPM não substitui o log com qty/stock; ele diz se vale abrir incidente.",
    )

    add_heading(doc, "13.5 Contraste com sucesso", 2)
    add_code(
        doc,
        "biz_event:sale_created          → sale_id, amount, plano\n"
        "biz_event:proposal_integrated   → outro domínio/serviço, mesmo fio de correlação",
    )
    add_para(
        doc,
        "Erro carrega error_type; sucesso carrega biz_event + IDs. A jornada pode mudar de "
        "domínio (venda → propostas) — a correlação une as metades.",
    )

    add_heading(doc, "13.6 Mapa sintoma → busca", 2)
    add_table(
        doc,
        ["Sintoma", "Vá para", "Busca"],
        [
            ["409 estoque", "Logs", "error_type:InsufficientStock"],
            ["Sessão de teste", "Logs", "ID de sessão / test_run_id"],
            ["Qual serviço falhou?", "Trace", "pelo trace_id"],
            ["Venda ok, proposta não?", "Logs", "sale_id / etapa integrate"],
            ["Erro subindo no time?", "RED", "serviço + operação"],
        ],
    )

    # ── 14 ──
    add_heading(doc, "14. Contrato de telemetria", 1)
    add_para(doc, "Antes de “instalar ferramenta X”, documente o contrato.")

    add_heading(doc, "14.1 Headers de entrada (cliente → API)", 2)
    add_table(
        doc,
        ["Header", "Obrigatório", "Descrição"],
        [
            ["X-Request-Id", "Sim (ou gerar no server)", "ID da chamada"],
            ["traceparent", "Ideal (W3C)", "Continuidade do trace"],
            ["ID de sessão / test run", "Recomendado em teste/jornada", "Várias calls"],
            ["X-Journey-Step", "Sim na jornada", "Etapa estável"],
            ["Staff/user", "Quando autenticado", "Operador"],
        ],
    )

    add_heading(doc, "14.2 Headers de saída", 2)
    add_para(doc, "Sempre: X-Request-Id + X-Trace-Id e/ou traceparent — inclusive em 4xx/5xx.")

    add_heading(doc, "14.3 Campos de log", 2)
    add_bullet(doc, "Identidade: serviço/domínio, trace_id, span_id, request_id")
    add_bullet(doc, "HTTP: method, path, route template, status, duration_ms")
    add_bullet(doc, "Negócio: biz_event, meio de pagamento, amount, item_id, …")
    add_bullet(doc, "Falha: error_type")

    add_heading(doc, "14.4 Etapas de jornada (exemplo lab)", 2)
    for s in [
        "sale.login", "sale.validate_customer", "sale.check_credit", "sale.search_product",
        "sale.check_stock", "sale.list_plans", "sale.simulate_plan", "sale.create_presale",
        "sale.list_proposals", "sale.integrate_proposal", "sale.create_cp",
    ]:
        add_bullet(doc, s)

    add_heading(doc, "14.5 Erros tipados (exemplo)", 2)
    for e in [
        "ValidationError", "Unauthorized", "ProductNotFound / CustomerNotFound",
        "InsufficientStock", "CreditNotApproved / CreditLimitExceeded / LimitsNotFound",
        "DbPoolTimeout / DbOperationalError",
    ]:
        add_bullet(doc, e)

    # ── 15 ──
    add_heading(doc, "15. Implementação do zero", 1)
    add_para(doc, "Ordem recomendada — pule etapas e o dashboard vira decoração.")
    add_numbered(doc, "Modelar a jornada — nomes estáveis de etapa (senão não há funil).")
    add_numbered(doc, "App + dependências com falha real — erros de negócio como HTTP + JSON tipado.")
    add_numbered(doc, "Logs estruturados — campos do contrato (se não é campo, não filtra bem).")
    add_numbered(doc, "Traces OTel/APM — HTTP (+ DB); export OTLP; spanmetrics/RED.")
    add_numbered(doc, "Correlação no cliente/BFF — headers; deep links só em ops/backoffice se quiser.")
    add_numbered(doc, "Dashboards com intenção — uma pergunta de negócio por widget.")
    add_numbered(doc, "Carga controlada — mix realista + alguns 4xx.")
    add_numbered(doc, "Mesmo contrato no deploy — nomes estáveis entre ambientes.")

    # ── 16 ──
    add_heading(doc, "16. Dashboards com intenção", 1)
    add_bullet(doc, "Negócio / Overview — KPIs, funil, mix, erros tipados, GMV.")
    add_bullet(doc, "Por domínio / funcionalidade — endpoints daquele bounded context + performance.")
    add_para(
        doc,
        "Pergunta guia ao criar widget: “Que decisão alguém toma ao ver este número?” "
        "Se não houver resposta, não crie o widget.",
        italic=True,
    )

    # ── 17 ──
    add_heading(doc, "17. Do laboratório ao sistema real", 1)
    add_para(
        doc,
        "Não é obrigatório copiar Graylog + Jaeger + Prometheus. É obrigatório copiar responsabilidades:",
    )
    add_table(
        doc,
        ["Responsabilidade", "Exemplos reais"],
        [
            ["Logs + error_type / biz_event", "Datadog Logs, Elastic, CloudWatch, Loki"],
            ["Trace + trace_id", "Datadog APM, New Relic, Tempo, Elastic APM"],
            ["RED por domínio", "Service pages, Prometheus + Grafana"],
            ["Deep link ops", "Backoffice / runbook / tip de suporte"],
            ["Carga / synthetic", "k6, JMeter, synthetics"],
        ],
    )
    add_para(doc, "Sugestão de adoção no produto real:", bold=True)
    add_numbered(doc, "Contrato de telemetria (1–2 páginas) com Produto + Engenharia.")
    add_numbered(doc, "Começar pela jornada crítica (ex.: pré-venda), não pelo monolito inteiro.")
    add_numbered(doc, "Logs JSON + trace_id no APM já existente.")
    add_numbered(doc, "Alertas por taxa de error_type (negócio), não só CPU.")
    add_numbered(doc, "SLOs simples: p95 da operação crítica; taxa de 5xx; error_type operacionais se fizer sentido.")

    # ── 18 ──
    add_heading(doc, "18. Maturidade", 1)
    add_heading(doc, "Curto prazo", 2)
    add_bullet(doc, "Quebrar de propósito e narrar o veredito em uma frase (playbook).")
    add_bullet(doc, "Criar um widget para uma única pergunta de negócio.")
    add_bullet(doc, "Desenhar o funil do sistema real com nomes de step.")
    add_heading(doc, "Médio prazo", 2)
    add_bullet(doc, "Contrato de telemetria oficial.")
    add_bullet(doc, "Instrumentação da jornada crítica.")
    add_bullet(doc, "Alertas e runbooks ligados a error_type / SLO.")
    add_heading(doc, "Longo prazo", 2)
    add_bullet(doc, "OpenTelemetry como padrão de vendor.")
    add_bullet(doc, "Sampling inteligente (100% erros + amostra de sucesso).")
    add_bullet(doc, "Exemplars (métrica → exemplo de trace).")
    add_bullet(doc, "Separar telemetria de produto (funil/GMV) da de plataforma (k8s/CPU).")

    # ── 19 ──
    add_heading(doc, "19. Anti-padrões", 1)
    for b in [
        "Dezenas de dashboards sem jornada definida.",
        "Log só em texto livre (“erro ao processar”).",
        "Só métricas de infra (“CPU alta”) sem ligar à operação de negócio.",
        "HTTP 4xx/5xx sem error_type — todos os conflitos parecem iguais.",
        "Depender de UI didática no app do usuário para achar trace_id.",
        "Trace sem os mesmos campos críticos no log (ou o contrário).",
        "Deep links com 127.0.0.1 para usuário remoto.",
        "Assumir crypto.randomUUID em HTTP não seguro sem fallback.",
        "Widgets multi-série no Graylog com id ≠ effectiveName (células em branco) — tip de lab.",
        "Payload enorme na argv do curl no bootstrap (ARG_MAX) — use arquivo/stdin.",
        "Estoque/crédito de lab que esgotam e invalidam demos — seed alto + restock/reset.",
    ]:
        add_bullet(doc, b)

    # ── 20 ──
    add_heading(doc, "20. Critério de prontidão", 1)
    add_para(doc, "Você está operacional quando, em qualquer ambiente, consegue:", bold=True)
    add_numbered(doc, "A partir de um sintoma, achar um log classificado.")
    add_numbered(doc, "Obter trace_id (direto ou via request_id / ID de negócio).")
    add_numbered(doc, "Abrir o trace e apontar o span culpado ou o gargalo.")
    add_numbered(doc, "Dizer se é isolado ou tendência (RED).")
    add_numbered(doc, "Explicar com campos de negócio, não só com status HTTP.")
    add_para(
        doc,
        "Ferramentas mudam; o método não — identidade compartilhada, classificação estável, "
        "e a ordem métrica → trace → log (ou o caminho inverso a partir do ticket).",
        italic=True,
    )

    # ── Anexos ──
    doc.add_page_break()
    add_heading(doc, "Anexo A — Queries úteis", 1)
    add_para(doc, "Adapte a sintaxe ao seu indexador (Graylog, Lucene, LogQL, etc.).")
    add_code(
        doc,
        "error_type:InsufficientStock\n"
        "error_type:CreditNotApproved\n"
        "biz_event:sale_created\n"
        "biz_event:proposal_integrated\n"
        "journey_step:sale.create_presale AND http_status:>=400\n"
        "duration_ms:>500 AND http_path:*CreatePreSales*\n"
        "trace_id:SEU_TRACE_ID\n"
        "request_id:\"SEU_REQUEST_ID\"\n"
        "test_run_id:\"ui-SEU_RUN\"",
    )

    add_heading(doc, "Anexo B — Arquitetura de referência (lab)", 1)
    add_para(
        doc,
        "Estratégia em uma frase: instrumentar a jornada como cadeia correlacionada; "
        "expor logs, traces e métricas RED; dashboards orientados a perguntas; "
        "playbook que classifica antes de otimizar.",
        italic=True,
    )
    add_code(
        doc,
        "Cliente (UI / curl / JMeter)\n"
        "   │  X-Request-Id · X-Test-Run-Id · X-Journey-Step · X-Staff-Id\n"
        "   ▼\n"
        "App (APIs da jornada)\n"
        "   ├── Postgres (estoque, crédito, pré-venda, propostas)\n"
        "   ├── Logs GELF ──► Graylog\n"
        "   └── OTel ─► Collector ─► Jaeger\n"
        "                    └── spanmetrics ─► Prometheus ◄─ Monitor (SPM)",
    )
    add_para(
        doc,
        "A aplicação do lab pode ser um único processo emitindo lab_domain / spm_service "
        "(assistente-venda, assistente-estoque, …) para simular bounded contexts. "
        "No sistema real, cada microserviço ocuparia naturalmente essa fronteira.",
    )

    add_heading(doc, "Anexo C — URLs e credenciais de treino", 1)
    add_table(
        doc,
        ["Serviço", "URL típica"],
        [
            ["Assistente UI", "http://leandrofdx.duckdns.org:8080"],
            ["Graylog", "https://leandrofdx.duckdns.org/"],
            ["Jaeger", "http://leandrofdx.duckdns.org:16686"],
            ["Jaeger SPM", "http://leandrofdx.duckdns.org:16686/monitor"],
        ],
    )
    add_para(doc, "Login Assistente (lab): conforme ambiente do time (ex.: vendedor1 / lab123).")
    add_para(doc, "Login Graylog (lab): admin / admin — trocar em qualquer ambiente compartilhado.")
    add_para(doc, "Restock / reset de crédito (lab):")
    add_code(
        doc,
        'curl -X POST http://HOST:8080/Stock/api/Stock/Restock -H \'Content-Type: application/json\' -d \'{"force":true}\'\n'
        "curl -X POST http://HOST:8080/Customer/api/ResetCreditLimit -H 'Content-Type: application/json' -d '{}'",
    )

    doc.add_paragraph()
    end = doc.add_paragraph()
    end.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = end.add_run(
        "— Fim do manual —\n\n"
        "Consolidado a partir do material de capacitação validado em chat.\n"
        "Use como referência de implementação e como material de treinamento.\n"
        "O laboratório ilustra; o método é o produto."
    )
    set_run_font(r, size=10, italic=True, color=MUTED)

    doc.save(OUT)
    print(f"OK: {OUT}")


if __name__ == "__main__":
    build()
